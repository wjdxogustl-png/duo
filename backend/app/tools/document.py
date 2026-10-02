"""한국어교육 신청서 초안 생성.

팀이 만든 템플릿(data/templates/korean_class_application.docx)의 {{자리표시자}}를 사용자 값으로 채운다.
템플릿이 없으면 기본 템플릿을 만들어 둔다. 실제 기관 제출은 하지 않는다.
"""
from datetime import datetime

from docx import Document

from .. import memory
from ..memory import DATA_DIR, OUTPUT_DIR

TEMPLATE_PATH = DATA_DIR / "templates" / "korean_class_application.docx"

FIELDS = {
    "name": "성명",
    "phone": "연락처",
    "region": "거주 시·군",
    "institution": "희망 기관",
    "preferred_time": "희망 시간대",
    "korean_level": "현재 한국어 수준",
}


def ensure_template() -> None:
    if TEMPLATE_PATH.exists():
        return
    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("한국어교육 수강 신청서 (초안)", level=1)
    table = doc.add_table(rows=len(FIELDS), cols=2)
    table.style = "Table Grid"
    for row, (key, label) in zip(table.rows, FIELDS.items()):
        row.cells[0].text = label
        row.cells[1].text = "{{" + key + "}}"
    doc.add_paragraph("작성일: {{date}}")
    doc.add_paragraph("※ 이 문서는 정착 도우미가 만든 초안입니다. 내용을 확인한 뒤 해당 기관에 직접 제출하세요.")
    doc.save(TEMPLATE_PATH)


def _replace(paragraph, values: dict) -> None:
    text = paragraph.text
    if "{{" not in text:
        return
    for k, v in values.items():
        text = text.replace("{{" + k + "}}", str(v))
    for run in paragraph.runs[1:]:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def fill_document(values: dict, out_name: str) -> str:
    """템플릿을 채워 OUTPUT_DIR/out_name 으로 저장하고 경로를 반환(테스트 대상)."""
    ensure_template()
    values = {**{k: "" for k in FIELDS}, **values, "date": datetime.now().strftime("%Y-%m-%d")}
    doc = Document(TEMPLATE_PATH)
    for p in doc.paragraphs:
        _replace(p, values)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace(p, values)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUTPUT_DIR / out_name
    doc.save(out)
    return str(out)


def generate_application_doc(name: str, phone: str, institution: str, preferred_time: str) -> dict:
    profile = memory.load()["profile"]
    level_names = {0: "거의 못함", 1: "기초", 2: "일상대화", 3: "능숙"}
    values = {
        "name": name,
        "phone": phone,
        "institution": institution,
        "preferred_time": preferred_time,
        "region": profile.get("region", ""),
        "korean_level": level_names.get(profile.get("korean_level"), ""),
    }
    fname = f"{memory._safe_id(memory.current_user.get())}_{datetime.now():%Y%m%d%H%M%S}.docx"
    fill_document(values, fname)
    return {"file": fname, "download_url": f"/api/files/{fname}", "filled": values}
