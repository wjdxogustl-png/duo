"""신청서 초안: 에이전트가 아는 정보로 지원사업 신청서의 기본 내용을 채워 두고, 사용자가 화면에서 고친다.

법·안전 원칙 (학생 프로젝트 수준에서 지킬 수 있는 것만 코드로 강제한다)
- 대신 제출하지 않는다. 초안을 보여 주고, 제출은 사용자가 기관에 직접 한다.
- 고유식별정보(외국인등록번호·주민등록번호·여권번호)와 계좌번호는 받지 않는다. 들어오면 지우고 알린다
  (개인정보 보호법: 고유식별정보는 법령 근거 없이 처리할 수 없음).
- 지어내지 않는다. 프로필에 없는 값은 비워 두고, 신청 동기는 사용자가 말한 사실로만 쓴다.
- 지울 수 있다. 초안은 사용자 상태에만 저장되고 '처음부터'(memory.reset)로 파일까지 지워진다.
기관의 공식 서식을 베끼지 않고, 대부분의 신청서에 공통으로 들어가는 항목만 쓴다.
"""
import re
import uuid
from datetime import datetime

from docx import Document

from .. import memory
from .programs import load_programs

MAX_DRAFTS = 5
MAX_TEXT = 600

LANG_LABELS = {"ko": "한국어", "en": "영어", "vi": "베트남어", "ja": "일본어", "zh": "중국어"}
LEVEL_LABELS = {0: "거의 못함", 1: "기초", 2: "일상 대화 가능", 3: "능숙"}

NOTICES = [
    "AI가 만든 초안이에요. 사실과 다른 내용이 없는지 꼭 확인하고 고쳐 주세요.",
    "외국인등록번호·여권번호·계좌번호는 여기에 적지 말고, 기관의 공식 신청서에만 직접 적으세요.",
    "제출은 본인이 기관 홈페이지·방문·전화로 직접 해요. 도우미는 대신 제출하지 않아요.",
    "입력한 내용은 초안을 만들기 위해 AI(Claude)에게 전달돼요. '처음부터'를 누르면 초안과 내려받은 파일이 모두 지워져요.",
]

# 외국인등록번호·주민등록번호(6자리-7자리), 여권번호(영문 1~2자 + 숫자 7~8자), 계좌번호처럼 보이는 긴 숫자열
# (\b 는 한글 옆에서 동작하지 않으므로 영문·숫자 경계를 직접 쓴다)
SENSITIVE = [
    ("외국인등록번호·주민등록번호", re.compile(r"(?<!\d)\d{6}\s*-?\s*[1-8]\d{6}(?!\d)")),
    ("여권번호", re.compile(r"(?<![A-Za-z0-9])[A-Za-z]{1,2}\d{7,8}(?!\d)")),
    ("계좌번호", re.compile(r"(?<![\d-])\d{2,6}-\d{2,6}-\d{2,8}(?:-\d{1,6})?(?![\d-])")),
]
PHONE = re.compile(r"(?<![\d-])0\d{1,2}-\d{3,4}-\d{4}(?![\d-])")


def scrub(text: str) -> tuple[str, list[str]]:
    """고유식별정보·계좌번호로 보이는 부분을 지운다(순수 함수). 전화번호는 남긴다."""
    found = []
    phones = PHONE.findall(text)
    masked = PHONE.sub("\0", text)  # 전화번호가 계좌번호 규칙에 걸리지 않게 잠시 가린다
    for label, pattern in SENSITIVE:
        if pattern.search(masked):
            found.append(label)
            masked = pattern.sub("[삭제됨]", masked)
    for p in phones:
        masked = masked.replace("\0", p, 1)
    return masked, found


def _program(program_id: str) -> dict:
    p = next((x for x in load_programs() if x["id"] == program_id), None)
    if p is None:
        raise ValueError(f"DB에 없는 사업: {program_id}. search_programs 결과의 id 를 넣으세요.")
    if p.get("is_sample"):
        raise ValueError(f"{program_id} 는 예시 데이터라 신청서를 만들 수 없어요.")
    return p


def _field(key, label, value="", hint="", multiline=False):
    return {"key": key, "label": label, "value": value if value is not None else "", "hint": hint, "multiline": multiline}


def build_fields(profile: dict, program: dict, preferred_time: str, motivation: str, requests: str) -> list[dict]:
    """신청서 칸을 만든다(순수 함수). 프로필에 없는 값은 비워 두고 무엇을 적을지 hint 로 알려 준다."""
    months = profile.get("months_in_korea")
    fields = [
        _field("name", "성명", profile.get("name"), "여권·외국인등록증과 같은 이름"),
        _field("phone", "연락처", "", "직접 입력 (기관이 연락할 번호)"),
        _field("region", "거주 시·군", profile.get("region")),
        _field("language", "편한 언어", LANG_LABELS.get(profile.get("language"), profile.get("language"))),
        _field("months_in_korea", "한국 거주 기간", f"{months}개월" if months is not None else ""),
        _field("korean_level", "한국어 수준", LEVEL_LABELS.get(profile.get("korean_level"), "")),
    ]
    if program.get("category") == "자녀교육" or "자녀" in program.get("target", ""):
        has = profile.get("has_children")
        fields.append(_field("children", "자녀", "" if has is None else ("있음" if has else "없음"),
                             "자녀 나이·학년 등 기관이 묻는 내용"))
    fields += [
        _field("preferred_time", "희망 일정·시간대", preferred_time, f"운영 일정: {program.get('schedule', '')}"),
        _field("motivation", "신청 동기", motivation, "왜 필요한지 내 말로 2~3문장", multiline=True),
        _field("requests", "요청 사항", requests, "통역 필요, 아이 동반 등", multiline=True),
    ]
    return fields


def draft_application(program_id: str, preferred_time: str = "", motivation: str = "", requests: str = "") -> dict:
    program = _program(program_id)
    state = memory.load()
    warnings = []
    cleaned = []
    for text in (preferred_time, motivation, requests):
        t, found = scrub((text or "")[:MAX_TEXT])
        cleaned.append(t)
        warnings += found
    fields = build_fields(state["profile"], program, *cleaned)
    draft = {
        "id": uuid.uuid4().hex[:8],
        "program": {k: program.get(k) for k in ("id", "name", "contact", "how_to_apply", "source_url", "address", "schedule")},
        "fields": fields,
        "notices": NOTICES,
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }
    state["drafts"] = (state.get("drafts") or [])[-(MAX_DRAFTS - 1):] + [draft]
    memory.save(state)
    empty = [f["label"] for f in fields if not f["value"]]
    return {
        "draft": draft,
        "empty_fields": empty,
        "removed_sensitive": sorted(set(warnings)),
        "next": "화면에 초안이 보인다. 사용자가 빈칸을 직접 채우고 고친 뒤, 내려받거나 복사해서 기관에 직접 제출한다.",
    }


def update_draft(draft_id: str, values: dict) -> dict:
    """화면에서 고친 값을 저장한다. 고유식별정보는 여기서도 지운다."""
    state = memory.load()
    draft = next((d for d in state.get("drafts") or [] if d["id"] == draft_id), None)
    if draft is None:
        raise KeyError(draft_id)
    removed = []
    for f in draft["fields"]:
        if f["key"] in values:
            f["value"], found = scrub(str(values[f["key"]] or "")[:MAX_TEXT])
            removed += found
    draft["updated_at"] = datetime.now().isoformat(timespec="seconds")
    memory.save(state)
    return {"draft": draft, "removed_sensitive": sorted(set(removed))}


def draft_docx(draft_id: str) -> dict:
    """저장된 초안을 Word 파일로 만든다. 기관 서식이 아니라 내용을 옮겨 적기 쉬운 표 한 장이다."""
    draft = next((d for d in memory.load().get("drafts") or [] if d["id"] == draft_id), None)
    if draft is None:
        raise KeyError(draft_id)
    doc = Document()
    doc.add_heading(f"{draft['program']['name']} 신청서 초안", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for f in draft["fields"]:
        row = table.add_row().cells
        row[0].text, row[1].text = f["label"], f["value"]
    p = draft["program"]
    doc.add_paragraph(f"신청 방법: {p.get('how_to_apply') or ''} / 문의: {p.get('contact') or ''}")
    if p.get("source_url"):
        doc.add_paragraph(f"안내 페이지: {p['source_url']}")
    doc.add_paragraph(f"작성일: {datetime.now():%Y-%m-%d}")
    for n in draft["notices"]:
        doc.add_paragraph("※ " + n)
    memory.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    name = f"{memory._safe_id(memory.current_user.get())}_draft_{draft_id}.docx"
    doc.save(memory.OUTPUT_DIR / name)
    return {"file": name, "download_url": f"/api/files/{name}"}
