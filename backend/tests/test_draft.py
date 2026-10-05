"""신청서 초안 검증. 정량 목표(신청서 입력값 반영 정확도 100%) 증빙에도 그대로 사용한다.

- 프로필 값이 지어내지 않고 그대로 들어가는지, 모르는 칸은 비는지
- 화면에서 고친 값이 Word 파일에 100% 그대로 들어가는지
- 고유식별정보(외국인등록번호·여권번호)·계좌번호는 지워지는지
- '처음부터'를 누르면 초안 파일까지 지워지는지
"""
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import memory
from app.main import app
from app.tools import draft

PROGRAM = {"id": "P1", "name": "김해 한국어 교실", "region": "김해", "category": "한국어교육", "target": "외국인",
           "schedule": "평일 저녁", "contact": "055-000-0000", "how_to_apply": "방문", "source_url": "https://example.go.kr"}
KIDS = {**PROGRAM, "id": "P2", "name": "자녀 언어발달", "category": "자녀교육", "target": "12세 이하 자녀"}


@pytest.fixture(autouse=True)
def tmp_data(tmp_path, monkeypatch):
    monkeypatch.setattr(memory, "USERS_DIR", tmp_path / "users")
    monkeypatch.setattr(memory, "OUTPUT_DIR", tmp_path / "output")
    monkeypatch.setattr(draft, "load_programs", lambda: [PROGRAM, KIDS, {**PROGRAM, "id": "S", "is_sample": True}])
    memory.current_user.set("test")
    s = memory.empty_state()
    s["profile"] = {"name": "Nguyen Van A", "region": "김해", "months_in_korea": 3, "korean_level": 0,
                    "language": "vi", "has_children": True}
    memory.save(s)


def values(d):
    return {f["key"]: f["value"] for f in d["fields"]}


def test_fills_only_known_facts():
    r = draft.draft_application("P1", "평일 저녁", "한국어를 배워 아이 학교 상담에 혼자 가고 싶어요.")
    v = values(r["draft"])
    assert (v["name"], v["region"], v["months_in_korea"], v["korean_level"], v["language"]) == \
        ("Nguyen Van A", "김해", "3개월", "거의 못함", "베트남어")
    assert v["phone"] == "" and v["requests"] == ""            # 모르는 칸은 지어내지 않고 비운다
    assert {"연락처", "요청 사항"} <= set(r["empty_fields"])
    assert "children" not in v                                  # 한국어교육 신청서엔 자녀 칸이 없다
    assert values(draft.draft_application("P2")["draft"])["children"] == "있음"
    assert [d["id"] for d in memory.load()["drafts"]][0] == r["draft"]["id"]   # 앞 초안도 남아 있다


def test_rejects_unknown_or_sample_program():
    with pytest.raises(ValueError):
        draft.draft_application("없는사업")
    with pytest.raises(ValueError):
        draft.draft_application("S")


@pytest.mark.parametrize("text, label", [
    ("등록번호 900101-5123456 입니다", "외국인등록번호·주민등록번호"),
    ("여권 M12345678입니다", "여권번호"),
    ("계좌 110-123-456789 로 주세요", "계좌번호"),
])
def test_scrub_removes_identifiers(text, label):
    out, found = draft.scrub(text)
    assert found == [label] and "[삭제됨]" in out


def test_scrub_keeps_phone_and_plain_text():
    assert draft.scrub("연락처 010-1234-5678, 평일 저녁 7시") == ("연락처 010-1234-5678, 평일 저녁 7시", [])


def test_edit_then_docx_reflects_values_exactly():
    d = draft.draft_application("P1")["draft"]
    edited = {"name": "Tran Thi B", "phone": "010-9876-5432", "preferred_time": "토요일 오전",
              "motivation": "아이 학교 상담에 혼자 가고 싶어요.", "requests": "베트남어 통역이 필요해요."}
    r = draft.update_draft(d["id"], edited)
    assert r["removed_sensitive"] == []
    path = memory.OUTPUT_DIR / draft.draft_docx(d["id"])["file"]
    cells = [c.text for row in Document(path).tables[0].rows for c in row.cells]
    for v in edited.values():                                   # 고친 값이 하나도 빠짐없이 그대로
        assert v in cells
    paragraphs = "\n".join(p.text for p in Document(path).paragraphs)
    assert "직접 해요" in paragraphs and "AI가 만든 초안" in paragraphs   # 안내 문구 포함


def test_edit_scrubs_identifiers():
    d = draft.draft_application("P1")["draft"]
    r = draft.update_draft(d["id"], {"requests": "외국인등록번호 900101-5123456"})
    assert values(r["draft"])["requests"] == "외국인등록번호 [삭제됨]"
    assert r["removed_sensitive"] == ["외국인등록번호·주민등록번호"]


def test_docx_filename_sanitized():
    memory.current_user.set("../../evil")
    memory.save({**memory.empty_state(), "profile": {}})
    d = draft.draft_application("P1")["draft"]
    r = draft.draft_docx(d["id"])
    assert r["file"].startswith("evil_") and (memory.OUTPUT_DIR / r["file"]).exists()


def test_api_save_docx_and_reset_deletes_files():
    client = TestClient(app)
    d = draft.draft_application("P1")["draft"]
    r = client.put(f"/api/drafts/test/{d['id']}", json={"values": {"phone": "010-1111-2222"}})
    assert r.status_code == 200 and values(r.json()["draft"])["phone"] == "010-1111-2222"
    f = client.post(f"/api/drafts/test/{d['id']}/docx").json()["file"]
    assert (memory.OUTPUT_DIR / f).exists()
    assert client.put("/api/drafts/test/nope", json={"values": {}}).status_code == 404
    assert client.get("/api/state/test").json()["drafts"][0]["id"] == d["id"]
    client.delete("/api/state/test")
    assert not (memory.OUTPUT_DIR / f).exists()                # 처음부터 → 파일까지 삭제
    assert client.get("/api/state/test").json()["drafts"] == []
