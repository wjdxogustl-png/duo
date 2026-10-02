"""규칙 기반 도구 검증. 정량 목표(신청서 반영 정확도 100%) 증빙에도 그대로 사용한다.

실행: (backend 폴더에서) python -m pytest -q
"""
from datetime import date

import pytest
from docx import Document

from app import memory
from app.tools import programs, risk
from app.tools.dday import compute_dday
from app.tools.document import fill_document
from app.tools.roadmap import generate_steps


@pytest.fixture(autouse=True)
def tmp_data(tmp_path, monkeypatch):
    monkeypatch.setattr(memory, "USERS_DIR", tmp_path / "users")
    monkeypatch.setattr(memory, "OUTPUT_DIR", tmp_path / "output")
    monkeypatch.setattr("app.tools.document.OUTPUT_DIR", tmp_path / "output")
    monkeypatch.setattr("app.tools.document.TEMPLATE_PATH", tmp_path / "tpl.docx")
    memory.current_user.set("test")


# --- 로드맵: 입력값별로 의도한 순서가 나오는지 ---

def test_roadmap_new_arrival_with_children():
    p = {"months_in_korea": 1, "korean_level": 0, "has_children": True, "job_status": "employed",
         "has_local_support": False}
    ids = [s["id"] for s in generate_steps(p)]
    assert ids == ["registration_check", "bank_phone", "health_insurance", "community",
                   "korean_class", "children_school", "law_counsel"]


def test_roadmap_settled_job_seeker():
    p = {"months_in_korea": 24, "korean_level": 2, "has_children": False, "job_status": "seeking",
         "workplace_issue": False, "has_local_support": True}
    ids = [s["id"] for s in generate_steps(p)]
    assert ids == ["job_support", "law_counsel"]


# --- 위험도 채점 ---

def test_risk_low_score_recommends_mentoring():
    p = {"korean_level": 0, "has_local_support": False, "workplace_issue": True,
         "knows_support_programs": False, "months_in_korea": 1}
    r = risk.compute_score(p)
    assert r["score"] == 0 and r["recommend_mentoring_first"] is True


def test_risk_high_score():
    p = {"korean_level": 3, "has_local_support": True, "workplace_issue": False,
         "knows_support_programs": True, "months_in_korea": 18}
    r = risk.compute_score(p)
    assert r["score"] == 100 and r["recommend_mentoring_first"] is False


def test_risk_threshold_boundary():
    # 25(한국어3) + 20(직장 어려움 없음) + 6(3개월) = 51 → 기준(50) 이상
    p = {"korean_level": 3, "has_local_support": False, "workplace_issue": False,
         "knows_support_programs": False, "months_in_korea": 3}
    assert risk.compute_score(p)["score"] == 51
    assert risk.compute_score(p)["recommend_mentoring_first"] is False


def test_risk_incomplete_profile_not_judged():
    # 필수 항목만 채우고 채점용 항목(has_local_support 등)이 빠진 프로필
    p = {"korean_level": 1, "months_in_korea": 2}
    r = risk.compute_score(p)
    assert r["complete"] is False
    assert r["recommend_mentoring_first"] is None
    assert set(r["missing"]) == {"has_local_support", "workplace_issue", "knows_support_programs"}


# --- 대화 기록 정리: user 로 시작, 역할 교대, 빈 내용 제거 ---

def test_clean_history():
    from app.agent import clean_history
    h = [
        {"role": "assistant", "content": "브리핑"},   # 앞쪽 assistant → 버림
        {"role": "user", "content": "안녕"},
        {"role": "assistant", "content": ""},         # 빈 내용 → 건너뜀
        {"role": "assistant", "content": "답1"},
        {"role": "assistant", "content": "브리핑2"},  # 연속 assistant → 합침
        {"role": "user", "content": "질문"},
    ]
    assert clean_history(h) == [
        {"role": "user", "content": "안녕"},
        {"role": "assistant", "content": "답1\n\n브리핑2"},
        {"role": "user", "content": "질문"},
    ]


# --- 능동 브리핑: 6시간 안에는 다시 하지 않고, force=True 면 한다 ---

def test_briefing_interval(monkeypatch):
    from app import agent
    calls = []
    monkeypatch.setattr(agent, "run", lambda *a, **kw: calls.append(1) or {"reply": "hi"})
    memory.save({**memory.empty_state(), "profile": {"region": "김해"}}, "test")

    assert agent.briefing("test") is not None
    assert memory.load("test")["last_briefing_at"] is not None
    assert agent.briefing("test") is None
    assert agent.briefing("test", force=True) is not None
    assert len(calls) == 2


# --- D-day ---

@pytest.mark.parametrize("end, label", [
    ("2026-11-01", "D-30"),
    ("2026-10-02", "D-Day"),
    ("2026-09-30", "D+2"),
    ("2027-01-01", "D-91"),
])
def test_dday(end, label):
    assert compute_dday(end, today=date(2026, 10, 2))["label"] == label


def test_dday_leap_year():
    assert compute_dday("2028-03-01", today=date(2028, 2, 28))["days_left"] == 2


# --- 지원사업 검색 ---

def test_search_programs_by_region_and_category():
    r = programs.search_programs(region="김해", category="한국어교육")
    assert r["count"] >= 1
    assert all(p["category"] == "한국어교육" for p in r["programs"])


@pytest.mark.parametrize("region", ["김해시", "창원특례시", "경상남도"])
def test_search_programs_region_variants(region):
    assert programs.search_programs(region=region)["count"] >= 1


# --- 신청서: 입력값이 100% 반영되는지 ---

def test_document_fields_filled_exactly():
    values = {"name": "Nguyen Van A", "phone": "010-0000-0000", "region": "김해",
              "institution": "테스트 기관", "preferred_time": "평일 저녁", "korean_level": "기초"}
    path = fill_document(values, "t.docx")
    cells = [c.text for row in Document(path).tables[0].rows for c in row.cells]
    for v in values.values():
        assert v in cells
    full = "\n".join(p.text for p in Document(path).paragraphs) + "\n".join(cells)
    assert "{{" not in full


def test_document_filename_sanitized(tmp_path):
    from app.tools.document import generate_application_doc
    memory.current_user.set("../../evil")
    r = generate_application_doc("A", "010", "기관", "저녁")
    assert r["file"].startswith("evil_")
    assert (tmp_path / "output" / r["file"]).exists()
